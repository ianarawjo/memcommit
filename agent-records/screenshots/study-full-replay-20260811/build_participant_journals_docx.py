"""Build editable, image-embedded Word journals from the replay Markdown.

The Markdown files remain the source of narrative truth.  This builder turns
each numbered capture into one landscape page so a missing renderer-specific
relative image link cannot hide the primary evidence.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)
LIGHT_BLUE = "E8EEF5"
BLACK = RGBColor(28, 31, 35)


@dataclass(frozen=True)
class Entry:
    section: str
    number: int
    title: str
    image_name: str
    image_alt: str
    fields: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class Journal:
    title: str
    intro: tuple[str, ...]
    issues_heading: str
    issues: tuple[str, ...]
    entries: tuple[Entry, ...]


def _strip_inline_markdown(value: str) -> str:
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
    return value.strip()


def parse_journal(path: Path) -> Journal:
    text = path.read_text(encoding="utf-8")
    title_match = re.search(r"^# (.+)$", text, re.MULTILINE)
    issue_match = re.search(r"^## (.*(?:Issues|이슈).*)$", text, re.MULTILINE)
    first_entry = re.search(r"^### 001 —", text, re.MULTILINE)
    if not (title_match and issue_match and first_entry):
        raise ValueError(f"Unexpected journal structure: {path}")

    intro_block = text[title_match.end() : issue_match.start()].strip()
    intro = tuple(
        " ".join(line.strip() for line in paragraph.splitlines())
        for paragraph in re.split(r"\n\s*\n", intro_block)
        if paragraph.strip()
    )

    issue_end_match = re.search(r"^## ", text[issue_match.end() :], re.MULTILINE)
    issue_end = (
        issue_match.end() + issue_end_match.start()
        if issue_end_match
        else first_entry.start()
    )
    issue_block = text[issue_match.end() : issue_end]
    issues = tuple(
        _strip_inline_markdown(match.group(1))
        for match in re.finditer(r"^\d+\. (.+)$", issue_block, re.MULTILINE)
    )

    section = ""
    entries: list[Entry] = []
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("## ") and not line.startswith("### "):
            section = line[3:].strip()
            index += 1
            continue
        heading = re.match(r"^### (\d{3}) — (.+)$", line)
        if not heading:
            index += 1
            continue
        number = int(heading.group(1))
        title = heading.group(2).strip()
        image_name = ""
        image_alt = ""
        fields: list[tuple[str, str]] = []
        index += 1
        while index < len(lines) and not lines[index].startswith(("## ", "### ")):
            image = re.match(
                r'^<img src="\./([^"]+\.png)" alt="([^"]*)" width="100%">$',
                lines[index],
            )
            if image:
                image_name, image_alt = image.groups()
            field = re.match(r"^- \*\*(.+?):\*\* (.+)$", lines[index])
            if field:
                fields.append((field.group(1), _strip_inline_markdown(field.group(2))))
            index += 1
        if not image_name or len(fields) != 3:
            raise ValueError(f"Incomplete entry {number:03d} in {path}")
        entries.append(
            Entry(
                section=section,
                number=number,
                title=title,
                image_name=image_name,
                image_alt=image_alt,
                fields=tuple(fields),
            )
        )

    if [entry.number for entry in entries] != list(range(1, 117)):
        raise ValueError(f"Journal does not contain ordered entries 001-116: {path}")
    return Journal(
        title=title_match.group(1).strip(),
        intro=intro,
        issues_heading=issue_match.group(1).strip(),
        issues=issues,
        entries=tuple(entries),
    )


def set_run_font(
    run,
    *,
    size: float,
    color: RGBColor = BLACK,
    bold: bool = False,
    italic: bool = False,
    east_asia: str = "Arial Unicode MS",
) -> None:
    # One Unicode-capable family on every script slot prevents LibreOffice from
    # dropping Hangul when a macOS-only East Asian fallback is unavailable.
    run.font.name = "Arial Unicode MS"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8, color=MUTED)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def set_image_alt(inline_shape, alt: str, title: str) -> None:
    doc_property = inline_shape._inline.docPr
    doc_property.set("descr", alt)
    doc_property.set("title", title)


def configure_document(doc: Document, language: str) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.header_distance = Inches(0.20)
    section.footer_distance = Inches(0.20)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    numbered = styles["List Number"]
    numbered.font.name = "Arial Unicode MS"
    numbered._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    numbered._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    numbered._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    numbered.font.size = Pt(9.5)
    numbered.paragraph_format.left_indent = Inches(0.375)
    numbered.paragraph_format.first_line_indent = Inches(-0.188)
    numbered.paragraph_format.space_after = Pt(4)
    numbered.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 28, DARK_BLUE, 0, 8),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_paragraph.add_run(
        "STUDY REPLAY · PARTICIPANT JOURNAL"
        if language == "en"
        else "STUDY 재현 · 참가자 저널"
    )
    set_run_font(header_run, size=8, color=MUTED, bold=True)
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc: Document, journal: Journal, language: str) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(34)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("USER STUDY EVIDENCE JOURNAL" if language == "en" else "USER STUDY 증거 저널")
    set_run_font(run, size=10, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    run = title.add_run(journal.title)
    set_run_font(run, size=28, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    text = (
        "116 ordered terminal captures · action · visible outcome · issue"
        if language == "en"
        else "터미널 캡처 116개 · 조작 · 확인된 화면 · 이슈"
    )
    run = subtitle.add_run(text)
    set_run_font(run, size=13, color=MUTED)

    values = (
        ("RUN", "study-snapshot-replay-20260811"),
        ("EVIDENCE", "180 × 52 color PTY"),
        ("FORMAT", "Editable Word · embedded images"),
    )
    if language == "ko":
        values = (
            ("실행", "study-snapshot-replay-20260811"),
            ("증거", "180 × 52 컬러 PTY"),
            ("형식", "수정 가능한 Word · 이미지 포함"),
        )
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(0)
    metadata.paragraph_format.space_after = Pt(16)
    metadata.paragraph_format.line_spacing = 1.25
    shade_paragraph(metadata, LIGHT_BLUE)
    for position, (label, value) in enumerate(values):
        if position:
            separator = metadata.add_run("    ·    ")
            set_run_font(separator, size=9, color=MUTED)
        label_run = metadata.add_run(f"{label} ")
        set_run_font(label_run, size=8, color=BLUE, bold=True)
        value_run = metadata.add_run(value)
        set_run_font(value_run, size=9, color=BLACK)

    for paragraph_text in journal.intro:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(_strip_inline_markdown(paragraph_text))
        set_run_font(run, size=10, color=BLACK)


def add_issue_page(doc: Document, journal: Journal) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_before = Pt(4)
    run = heading.add_run(journal.issues_heading)
    set_run_font(run, size=16, color=BLUE, bold=True)

    for issue in journal.issues:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(issue)
        set_run_font(run, size=9.5, color=BLACK)


def add_entry_page(doc: Document, entry: Entry, language: str) -> None:
    doc.add_page_break()

    section = doc.add_paragraph()
    section.paragraph_format.space_before = Pt(0)
    section.paragraph_format.space_after = Pt(1)
    section.paragraph_format.keep_with_next = True
    run = section.add_run(entry.section.upper())
    set_run_font(run, size=8, color=BLUE, bold=True)

    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run(f"{entry.number:03d} — {entry.title}")
    set_run_font(run, size=12, color=DARK_BLUE, bold=True)

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(0)
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.paragraph_format.keep_with_next = True
    image_path = ROOT / entry.image_name
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    run = image_paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(9.25))
    set_image_alt(
        inline_shape,
        entry.image_alt,
        f"Capture {entry.number:03d}: {entry.title}",
    )

    for label, value in entry.fields:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1.2)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_together = True
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=8.2, color=BLUE, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=8.2, color=BLACK)


def build(language: str) -> Path:
    source = ROOT / f"PARTICIPANT-JOURNAL.{language}.md"
    output = ROOT / f"PARTICIPANT-JOURNAL.{language}.docx"
    journal = parse_journal(source)
    document = Document()
    configure_document(document, language)
    add_cover(document, journal, language)
    add_issue_page(document, journal)
    for entry in journal.entries:
        add_entry_page(document, entry, language)
    properties = document.core_properties
    properties.title = journal.title
    properties.subject = "Observer-authored journal of the full Study replay"
    properties.author = "Memcommit Research Team"
    properties.keywords = "user study, participant journal, terminal evidence, memcommit"
    document.save(output)
    return output


if __name__ == "__main__":
    for code in ("en", "ko"):
        path = build(code)
        print(path)
